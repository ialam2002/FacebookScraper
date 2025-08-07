"""
Tab for scraping people who liked posts from profile URLs.
"""
import customtkinter as ctk
from tkinter import messagebox, filedialog, scrolledtext
import threading
import json
import os
from scraper.post_scraper import FacebookPostsScraper

class PostScraperTab:
    def __init__(self, parent, app):
        self.app = app
        self.frame = ctk.CTkFrame(parent)
        self.frame.pack(fill="both", expand=True, padx=10, pady=10)
        self.post_scraper = None
        self.scraping_active = False
        self.loaded_json_results = None  # To store results from loaded JSON
        self.build_post_scraper_tab()

    def build_post_scraper_tab(self):
        frame = self.frame
        header = ctk.CTkLabel(frame, text="Post Likes Scraper", font=ctk.CTkFont(size=16, weight="bold"))
        header.pack(pady=(0, 15))

        # Instructions
        instructions = ctk.CTkLabel(
            frame, 
            text="Enter profile URLs to scrape people who liked their posts:",
            font=ctk.CTkFont(size=12)
        )
        instructions.pack(anchor="w", padx=10, pady=(0, 10))

        # URLs input
        urls_label = ctk.CTkLabel(frame, text="Profile URLs (one per line):")
        urls_label.pack(anchor="w", padx=10, pady=(0, 5))

        self.profile_urls_text = scrolledtext.ScrolledText(frame, height=8, wrap="word", font=("Consolas", 10))
        self.profile_urls_text.pack(fill="x", padx=10, pady=(0, 15))

        # Parameters frame
        params_frame = ctk.CTkFrame(frame)
        params_frame.pack(fill="x", padx=10, pady=10)

        # Cap toggle for max posts per profile
        self.cap_enabled_var = ctk.BooleanVar(value=False)
        self.cap_checkbox = ctk.CTkCheckBox(params_frame, text="Enable Post Cap", variable=self.cap_enabled_var, onvalue=True, offvalue=False)
        self.cap_checkbox.grid(row=0, column=0, padx=5, pady=5, sticky="w")

        posts_label = ctk.CTkLabel(params_frame, text="Max Posts per Profile:")
        posts_label.grid(row=0, column=1, padx=5, pady=5, sticky="w")
        self.max_posts_entry = ctk.CTkEntry(params_frame, width=80)
        self.max_posts_entry.insert(0, "10")
        self.max_posts_entry.grid(row=0, column=2, padx=5, pady=5, sticky="w")
        # Disable max_posts_entry unless cap is enabled
        def toggle_max_posts_entry():
            state = "normal" if self.cap_enabled_var.get() else "disabled"
            self.max_posts_entry.configure(state=state)
        self.cap_enabled_var.trace_add('write', lambda *args: toggle_max_posts_entry())
        toggle_max_posts_entry()

        # Output file selection
        output_label = ctk.CTkLabel(frame, text="Output File (Optional - leave blank to save in memory only):")
        output_label.pack(anchor="w", padx=10, pady=(10, 5))

        output_frame = ctk.CTkFrame(frame, fg_color="transparent")
        output_frame.pack(fill="x", padx=10, pady=5)

        self.output_file_entry = ctk.CTkEntry(output_frame, width=400)
        self.output_file_entry.pack(side="left", padx=(0, 5))

        browse_btn = ctk.CTkButton(output_frame, text="Browse", command=self.browse_output_file, width=80)
        browse_btn.pack(side="left")

        # Control buttons
        buttons_frame = ctk.CTkFrame(frame, fg_color="transparent")
        buttons_frame.pack(fill="x", padx=10, pady=20)

        self.start_btn = ctk.CTkButton(
            buttons_frame, 
            text="Start Scraping", 
            command=self.start_scraping,
            fg_color="green",
            hover_color="darkgreen"
        )
        self.start_btn.pack(side="left", padx=(0, 10))

        self.stop_btn = ctk.CTkButton(
            buttons_frame, 
            text="Stop Scraping", 
            command=self.stop_scraping,
            fg_color="red",
            hover_color="darkred",
            state="disabled"
        )
        self.stop_btn.pack(side="left", padx=(0, 10))

        self.export_excel_btn = ctk.CTkButton(
            buttons_frame,
            text="Export Likes to Excel",
            command=self.export_likes_to_excel,
            fg_color="#1d6f42",
            hover_color="#14532d"
        )
        self.export_excel_btn.pack(side="left")

        # Progress label for status updates
        self.progress_label = ctk.CTkLabel(frame, text="", font=ctk.CTkFont(size=12, weight="bold"), text_color="#1d4f6f")
        self.progress_label.pack(fill="x", padx=10, pady=(0, 5))
        # Results text area (for status and loaded JSON summary)
        self.results_text = scrolledtext.ScrolledText(frame, height=8, wrap="word", font=("Consolas", 10))
        self.results_text.pack(fill="x", padx=10, pady=(0, 10))

        # Add button to load JSON results
        self.load_json_btn = ctk.CTkButton(
            buttons_frame,
            text="Load Likes JSON",
            command=self.load_likes_json,
            fg_color="#1d4f6f",
            hover_color="#14526d"
        )
        self.load_json_btn.pack(side="left", padx=(10, 0))

    def load_likes_json(self):
        """Load post likes data from a JSON file and display summary."""
        import json
        from tkinter import filedialog, messagebox
        file_path = filedialog.askopenfilename(
            defaultextension=".json",
            filetypes=[("JSON files", "*.json"), ("All files", "*.*")],
            title="Select Likes JSON File"
        )
        if not file_path:
            return
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                all_results = json.load(f)
        except Exception as e:
            messagebox.showerror("Error", f"Could not load JSON file: {e}")
            return
        # Store loaded results for export
        self.loaded_json_results = all_results
        # Display summary in results_text
        self.results_text.delete("1.0", "end")
        summary_lines = []
        for profile in all_results:
            name = profile.get("profile_name", "Unknown")
            url = profile.get("profile_url", "")
            post_count = len(profile.get("post_likes", []))
            summary_lines.append(f"{name} ({url}) - {post_count} posts scraped")
        self.results_text.insert("end", "Loaded JSON file:\n" + "\n".join(summary_lines) + "\n")
        self.results_text.see("end")

    def export_likes_to_excel(self):
        """Export deduplicated post likes to Excel."""
        import openpyxl
        from openpyxl.utils import get_column_letter
        from tkinter import filedialog, messagebox
        # Ask for file
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            title="Save Likes Excel File"
        )
        if not file_path:
            return
        # Use loaded JSON if available, else load from output file
        all_results = getattr(self, "loaded_json_results", None)
        if all_results is None:
            output_file = self.output_file_entry.get().strip()
            if output_file and os.path.exists(output_file):
                try:
                    with open(output_file, 'r', encoding='utf-8') as f:
                        all_results = json.load(f)
                except Exception as e:
                    messagebox.showerror("Error", f"Could not load JSON results: {e}")
                    return
        if not all_results:
            messagebox.showerror("Error", "No results to export. Please scrape first, load a JSON file, or select a valid output file.")
            return
        # Build deduplicated list: (scraped_profile_name, scraped_profile_url, liker_name, liker_profile_url)
        deduped = set()
        rows = []
        for profile in all_results:
            scraped_name = profile.get('profile_name', 'Unknown')
            scraped_url = profile.get('profile_url', '')
            for post in profile.get('post_likes', []):
                for liker in post.get('liked_by', []):
                    liker_name = liker.get('name', '')
                    liker_url = liker.get('profile_url', '')
                    key = (scraped_name, scraped_url, liker_name, liker_url)
                    if key not in deduped:
                        rows.append([scraped_name, scraped_url, liker_name, liker_url])
                        deduped.add(key)
        # Write to Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Post Likes"
        ws.append(["Scraped Profile Name", "Scraped Profile URL", "Liker Name", "Liker Profile URL"])
        for row in rows:
            ws.append(row)
        # Auto-size columns
        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 2, 50)
        try:
            wb.save(file_path)
            messagebox.showinfo("Export Successful", f"Likes exported to {file_path}")
        except Exception as e:
            messagebox.showerror("Export Failed", f"Could not save Excel file: {e}")

        # (No widget creation here; only in build_post_scraper_tab)

    def browse_output_file(self):
        """Browse for output file location."""
        filename = filedialog.asksaveasfilename(
            defaultextension=".json",
            filetypes=[("JSON files", "*.json"), ("All files", "*.*")],
            title="Save post likes data as..."
        )
        if filename:
            self.output_file_entry.delete(0, "end")
            self.output_file_entry.insert(0, filename)

    def start_scraping(self):
        """Start the post likes scraping process."""
        # Get URLs from text area
        urls_text = self.profile_urls_text.get("1.0", "end-1c").strip()
        if not urls_text:
            messagebox.showerror("Error", "Please enter at least one profile URL.")
            return

        # Parse URLs (handle both comma-separated and newline-separated)
        raw_urls = []
        for line in urls_text.split('\n'):
            raw_urls.extend([url.strip() for url in line.split(',') if url.strip()])
        # Only keep valid Facebook URLs
        urls = [u for u in raw_urls if 'facebook.com' in u]

        if not urls:
            messagebox.showerror("Error", "No valid Facebook URLs found.")
            return

        # Get max posts
        cap_enabled = self.cap_enabled_var.get() if hasattr(self, 'cap_enabled_var') else False
        if cap_enabled:
            try:
                max_posts = int(self.max_posts_entry.get())
                if max_posts <= 0:
                    raise ValueError()
            except ValueError:
                messagebox.showerror("Error", "Max posts must be a positive integer.")
                return
        else:
            max_posts = 1000000  # Effectively unlimited

        # Get output file
        output_file = self.output_file_entry.get().strip()
        if not output_file:
            output_file = None  # Allow None for memory-only storage

        # Check if logged in
        if not self.app.scraper_controller.is_logged_in():
            messagebox.showerror("Error", "Please log in first in the Configuration tab.")
            return

        # Start scraping in a separate thread
        self.scraping_active = True
        self.start_btn.configure(state="disabled")
        self.stop_btn.configure(state="normal")
        
        scraping_thread = threading.Thread(
            target=self.scrape_post_likes,
            args=(urls, max_posts, output_file),
            daemon=True
        )
        scraping_thread.start()

    def stop_scraping(self):
        """Stop the scraping process."""
        self.scraping_active = False
        self.start_btn.configure(state="normal")
        self.stop_btn.configure(state="disabled")
        self.update_progress("Scraping stopped by user.")

    def scrape_post_likes(self, urls, max_posts, output_file):
        """Perform the actual post likes scraping."""
        try:
            # Get the driver from the main scraper controller
            driver = self.app.scraper_controller.get_driver()
            if not driver:
                self.update_progress("Error: No active browser session found.")
                return

            # Create post scraper instance
            self.post_scraper = FacebookPostsScraper(driver)
            
            all_results = []
            
            for i, url in enumerate(urls):
                if not self.scraping_active:
                    break
                    
                self.update_progress(f"Processing profile {i+1}/{len(urls)}: {url}")
                
                try:
                    result = self.post_scraper.scrape_posts(url, max_posts)
                    if result:
                        all_results.append(result)
                        self.update_progress(f"✓ Scraped {len(result.get('post_likes', []))} posts from {result.get('profile_name', 'Unknown')}")
                    else:
                        self.update_progress(f"✗ Failed to scrape {url}")
                        
                except Exception as e:
                    self.update_progress(f"✗ Error scraping {url}: {str(e)}")
                    continue

            # Save results
            if all_results:
                # Store results in memory for access by other tabs
                self.loaded_json_results = all_results
                
                if output_file:
                    try:
                        with open(output_file, 'w', encoding='utf-8') as f:
                            json.dump(all_results, f, indent=2, ensure_ascii=False)
                        
                        self.update_progress(f"✓ Results saved to {output_file}")
                        self.update_progress(f"✓ Scraping completed! Total profiles: {len(all_results)}")
                        
                    except Exception as e:
                        self.update_progress(f"✗ Error saving results: {str(e)}")
                        self.update_progress(f"✓ Scraping completed! Total profiles: {len(all_results)} (data available for visualization)")
                else:
                    self.update_progress(f"✓ Scraping completed! Total profiles: {len(all_results)} (data saved in memory - available for visualization and export)")
            else:
                self.update_progress("✗ No data was scraped.")

        except Exception as e:
            self.update_progress(f"✗ Scraping error: {str(e)}")
        
        finally:
            # Reset UI state
            self.scraping_active = False
            self.start_btn.configure(state="normal")
            self.stop_btn.configure(state="disabled")

    def update_progress(self, message):
        """Update the progress display."""
        def update_ui():
            self.progress_label.configure(text=message)
            self.results_text.insert("end", f"{message}\n")
            self.results_text.see("end")
        
        # Schedule UI update on main thread
        self.app.after(0, update_ui)

    def clear_results(self):
        """Clear the results text area."""
        self.results_text.delete("1.0", "end")
